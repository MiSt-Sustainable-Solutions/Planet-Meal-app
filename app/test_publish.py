"""
A client sees what MiSt published, and nothing MiSt is still working on.

Until 13 Sep 2026 counting a file showed it to the client the same second, so there was no
way to look at figures before the client did. Now a client reads a published copy. This
checks the promise from both sides:

  * before anything is published, a client sees no figures, files or downloads at all
  * after publishing, they see exactly the figures MiSt saw when pressing the button
  * counting another file, a catalogue change, or a file being deleted afterwards changes
    NOTHING for them -- figures, files, downloads -- until MiSt publishes again
  * the admin bar says whether the client's copy still matches, and why not
  * "See what the client sees" shows the client's page, not the working one
  * a publish that fails leaves the client's copy exactly where it was
  * one client's publication is never another's

    python test_publish.py        (needs the catalogue API running)
"""
import io
import json
import os
import sys
import tempfile

_HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, _HERE)

_TMP = tempfile.mkdtemp(prefix="pp_publish_")
os.environ["MIST_APP_DB"] = os.path.join(_TMP, "publish_test.db")
os.environ["MIST_TENANT"] = "alpha"
os.environ["MIST_SECRET_KEY"] = "test-only-key-not-a-real-secret"

from fastapi.testclient import TestClient   # noqa: E402
from openpyxl import load_workbook          # noqa: E402

import auth        # noqa: E402
import catalogue   # noqa: E402
import db          # noqa: E402
import main        # noqa: E402
import publish     # noqa: E402
import selection   # noqa: E402
import store       # noqa: E402

FAILED = []


def P(ok, msg):
    print(("  PASS " if ok else "  FAIL ") + msg)
    if not ok:
        FAILED.append(msg)


if catalogue.health() is None:
    print("the catalogue API is not running -- start it and try again")
    sys.exit(2)

db.init()
auth.init()
auth.add_tenant("alpha", "Alpha University", "")
auth.add_tenant("beta", "Beta College", "")

# Alpha: one file counted (Jan-Jun, 100 kg a line), one held back (Jul-Sep, 300 kg a line).
# Beta: one file counted, with a volume nothing of Alpha's could be mistaken for.
FILES = [("upl-a1", "alpha", "Alpha spring.xlsx", range(1, 7), 100.0, 1),
         ("upl-a2", "alpha", "Alpha summer.xlsx", range(7, 10), 300.0, 0),
         ("upl-b1", "beta", "Beta year.xlsx", range(1, 7), 900.0, 1)]
con = db.connect()
for tenant in ("alpha", "beta"):
    store.upsert(con, "product", ["tenant", "artikelnr", "description", "category", "ean",
                                  "ean_he", "first_seen", "last_seen"],
                 [(tenant, "194072", "MEYERIJ VOLLE MELK", "ZUIVEL HOUDBAAR",
                   "8710401996797", "", "2025-01", "2025-09")],
                 conflict=["tenant", "artikelnr"])
for uid, tenant, name, months, kg, selected in FILES:
    for m in months:
        con.execute("INSERT INTO purchase_line (tenant, year, month, klantnr, restaurant, "
                    "city, artikelnr, aantal, omzet, kg, kg_known, quality, source_upload) "
                    "VALUES (?,2025,?,'K1','CANTEEN','X','194072',10,?,?,1,'complete',?)",
                    (tenant, m, kg * 5, kg, uid))
    periods = [f"2025-{m:02d}" for m in months]
    store.upsert(con, "upload",
                 ["id", "tenant", "filename", "stored_path", "adapter", "year", "uploaded_at",
                  "periods", "lines", "products", "spend_eur", "verdict", "report_json",
                  "committed_at", "commit_mode", "commit_note", "selected", "archived_at"],
                 [(uid, tenant, name, "", "sligro", 2025, "2025-10-01T00:00:00",
                   json.dumps(periods), len(periods), 1, kg * 5 * len(periods), "go",
                   json.dumps(dict(verdict="go", findings=[], summary={}, filename=name)),
                   None, None, None, selected, None)],
                 conflict=["id"])
con.commit()
con.close()
db.invalidate()

auth.create_user("mist", "admin-password-1", "admin", None, "MiSt")
auth.create_user("alphauser", "alpha-password-1", "client", "alpha", "Alpha University")
auth.create_user("betauser", "beta-password-1", "client", "beta", "Beta College")


def signed_in(username, password):
    c = TestClient(main.app, follow_redirects=False)
    assert c.post("/login", data={"username": username, "password": password}).status_code == 303
    return c


def kg(client, window=""):
    r = client.get("/api/analysis" + (f"?window={window}" if window else ""))
    return r.json()["headline"]["food_kg"] if r.status_code == 200 else None


def publish_as(admin, tenant):
    admin.post("/admin/viewing", data={"tenant": tenant, "back": "/"})
    r = admin.post("/admin/publish", data={"back": "/"})
    pid = publish.latest(tenant)["id"]
    return r, publish.wait(pid, timeout=600)


def xlsx_kg(data: bytes) -> float:
    ws = load_workbook(io.BytesIO(data), read_only=True)["summary"]
    return dict((r[0], r[1]) for r in ws.iter_rows(values_only=True))["food purchased (kg)"]


M = signed_in("mist", "admin-password-1")
A = signed_in("alphauser", "alpha-password-1")
B = signed_in("betauser", "beta-password-1")
M.post("/admin/viewing", data={"tenant": "alpha", "back": "/"})

print("=== before anything is published, a client sees nothing ===")
page = A.get("/").text
P("being prepared" in page, "the dashboard says their figures are being prepared")
P("CANTEEN" not in page.upper() and "Download Excel" not in page,
  "with no figures, restaurants or downloads on it")
P(kg(A) is None and A.get("/api/analysis").status_code == 404, "no figures by API either")
P("being prepared" in A.get("/data-health").text, "nor on Data health")
files = A.get("/files").text
P("Alpha spring.xlsx" not in files, "the Files page lists no file")
P(A.get("/export.xlsx").status_code == 404, "the Excel download is refused")
P(A.get("/files/upl-a1/lines.xlsx").status_code == 404, "and so is a file's lines")
P(A.get("/api/months").json()["rows"] == [], "and no months")

print("\n=== while MiSt sees the working figures, and is told the client sees nothing ===")
P(kg(M) == 600, f"the admin sees Alpha's counted figures ({kg(M)} kg)")
bar = M.get("/").text
P("Alpha University sees nothing yet." in bar, "the admin bar says Alpha sees nothing yet")
P("Publish to Alpha University" in bar, "and offers to publish")

print("\n=== nothing counted, nothing to publish ===")
selection.set_selected("upl-b1", "beta", False)
M.post("/admin/viewing", data={"tenant": "beta", "back": "/"})
bar = M.get("/files").text
P("Count a file first" in bar and "Publish to Beta College" not in bar,
  "with no file counted the button is not offered")
P(M.post("/admin/publish", data={"back": "/files"}).headers["location"].startswith(
    "/files?pub_error="), "and posting anyway is refused with a reason")
selection.set_selected("upl-b1", "beta", True)

print("\n=== publish ===")
r, pub = publish_as(M, "alpha")
P(r.status_code == 303, "the button starts it and returns straight away")
P(pub["status"] == "live", f"the publication finished ({pub['status']}: {pub.get('error')})")
P({w["key"] for w in pub["windows"]} == {"FY2025", "all"},
  f"it holds every period ({', '.join(w['key'] for w in pub['windows'])})")
P([f["upload_id"] for f in pub["files"]] == ["upl-a1"],
  "and the counted file, and only that one")

print("\n=== the client now sees exactly what MiSt saw ===")
P(kg(A) == 600, f"Alpha sees 600 kg ({kg(A)})")
page = A.get("/").text
P("Figures of " in page and "being prepared" not in page, "dated, and no longer 'being prepared'")
P("Alpha spring.xlsx" in page and 'name="file"' not in page,
  "the picker offers the published file on its own, and no ticking files together")
P(kg(A, "files:upl-a1") == 600, "choosing that file shows it")
P(kg(A, "files:upl-a2") == 600, "asking for the unpublished file by URL shows the default instead")
got = A.get("/export.xlsx?window=FY2025")
P(got.status_code == 200 and xlsx_kg(got.content) == 600,
  "the Excel download is the published workbook, with the same total")
P(A.get("/export.xlsx?window=files:upl-a2").status_code == 404,
  "an unpublished window cannot be downloaded")
P(A.get("/files/upl-a1/lines.xlsx").status_code == 200, "the published file's lines download")
P(A.get("/files/upl-a2/lines.xlsx").status_code == 404, "the held-back file's do not")
P(len(A.get("/api/months").json()["rows"]) == 6, "six months, as published")
run_id = A.get("/api/analysis").json().get("run_id")
P(A.get(f"/api/run/{run_id}").status_code == 200, "a published run id opens")
bar = M.get("/").text
P("sees the figures published on" in bar and "They match what you see." in bar,
  "the admin bar says the copy matches")
P("Publish to Alpha University" not in bar, "so there is nothing to publish")

print("\n=== MiSt keeps working; the client sees none of it ===")
selection.set_selected("upl-a2", "alpha", True)
P(kg(M) == 1500, f"counting the summer file moves the admin's figure ({kg(M)} kg)")
P(kg(A) == 600, f"and not the client's ({kg(A)} kg)")
P("Alpha summer.xlsx" not in A.get("/files").text, "the new file is not on their Files page")
P("Alpha summer.xlsx" not in A.get("/").text, "nor in their picker")
P(A.get("/files/upl-a2/lines.xlsx").status_code == 404, "nor downloadable")
unpublished_run = M.get("/api/analysis").json().get("run_id")
P(A.get(f"/api/run/{unpublished_run}").status_code == 404,
  "and MiSt's working run cannot be read by id, though it is the same client's")
bar = M.get("/").text
P("still sees the figures published on" in bar and "the counted files changed" in bar,
  "the admin bar says the copy is behind, and why")
P("Publish to Alpha University" in bar, "and offers to publish again")
P(A.get("/?refresh=1") is not None and kg(A) == 600, "a client's ?refresh=1 changes nothing")

print("\n=== a catalogue change does not reach the client either ===")
real_version, real_health = catalogue.version, catalogue.health
live_v = real_version()
moved = dict(live_v, version=live_v["version"] + 1, etag=f"cat-{live_v['version'] + 1}.moved")
# Both, because a page reads the version from the health check and scoring from version().
catalogue.version = lambda: moved
catalogue.health = lambda: dict(real_health(), catalogue=moved)
try:
    P("the catalogue or the way figures are calculated changed" in M.get("/").text,
      "the admin bar names a catalogue change as a reason")
    P(kg(A) == 600, "the client's figure is unchanged")
finally:
    catalogue.version, catalogue.health = real_version, real_health

print("\n=== see what the client sees ===")
M.post("/admin/as-client", data={"on": 1, "back": "/"})
page = M.get("/").text
P("You are seeing exactly what Alpha University sees" in page, "the bar says so")
P(kg(M) == 600, f"and the admin now reads the published 600 kg ({kg(M)})")
P('href="/catalogue"' not in page, "drawn as the client's page: no admin navigation")
P("Alpha summer.xlsx" not in M.get("/files").text, "the Files page is the client's too")
M.post("/admin/as-client", data={"on": 0, "back": "/"})
P(kg(M) == 1500, "and back to the working figures")

print("\n=== a publish that fails leaves the client's copy alone ===")
real_score = catalogue.score_lines


def _down(*a, **k):
    raise catalogue.CatalogueDown("could not reach the catalogue API (test)")


catalogue.score_lines = _down
try:
    _r, failed = publish_as(M, "alpha")
finally:
    catalogue.score_lines = real_score
P(failed["status"] == "failed" and "could not reach" in (failed["error"] or ""),
  f"the attempt is recorded as failed ({failed['error']})")
P(kg(A) == 600, "the client still sees the previous figures")
P(publish.live("alpha")["id"] == pub["id"], "the previous publication is still the live one")
con = db.connect()
left = con.execute("SELECT COUNT(*) FROM publication_file WHERE publication_id=?",
                   (failed["id"],)).fetchone()[0]
con.close()
P(left == 0, "and the failed attempt left no downloads behind")
P("The last publish did not finish" in M.get("/").text, "the admin bar says it did not finish")

print("\n=== publishing again brings the client up to date ===")
_r, pub2 = publish_as(M, "alpha")
P(pub2["status"] == "live", f"published ({pub2['status']}: {pub2.get('error')})")
P(kg(A) == 1500, f"Alpha now sees 1500 kg ({kg(A)})")
P("Alpha summer.xlsx" in A.get("/files").text, "and the summer file")
P(publish.get(pub["id"])["status"] == "replaced", "the old copy is recorded as replaced")
con = db.connect()
old_files = con.execute("SELECT COUNT(*) FROM publication_file WHERE publication_id=?",
                        (pub["id"],)).fetchone()[0]
con.close()
P(old_files == 0, "and its downloads are cleared, since nobody can reach them")

print("\n=== deleting a published file's lines does not reach the client ===")
selection.archive("upl-a2", "alpha")
selection.destroy("upl-a2", "alpha")
P(kg(M) == 600, "the admin's figure drops")
P(kg(A) == 1500 and A.get("/files/upl-a2/lines.xlsx").status_code == 200,
  "the client's published figures and downloads stay whole")

print("\n=== one client's publication is not another's ===")
P(kg(B) is None, "Beta, never published, still sees nothing")
P(B.get("/files/upl-a1/lines.xlsx").status_code == 404, "and cannot download Alpha's file")
_r, pubb = publish_as(M, "beta")
P(pubb["status"] == "live" and kg(B) == 5400, f"Beta's own publish shows Beta's 5400 kg ({kg(B)})")
P(kg(A) == 1500, "and leaves Alpha's untouched")
P(A.post("/admin/publish", data={"back": "/"}).status_code == 403,
  "a client cannot publish")
P(A.post("/admin/as-client", data={"on": 1}).status_code == 403,
  "nor use the preview switch")

print("\n=== a publish interrupted by a restart does not block the next ===")
con = db.connect()
con.execute("INSERT INTO publication (id, tenant, status, started_at) VALUES "
            "('ghost', 'beta', 'building', '2026-01-01T00:00:00')")
con.commit()
con.close()
P(publish.recover(startup=True) == 1, "startup marks it failed")
P(publish.get("ghost")["status"] == "failed", "with a reason")
P(kg(B) == 5400, "and Beta's figures are untouched")

print(f"\n{'ALL PASS' if not FAILED else str(len(FAILED)) + ' FAILED'}")
sys.exit(1 if FAILED else 0)
