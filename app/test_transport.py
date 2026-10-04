"""
Transport from the wholesaler's DC to the client, the way TU Delft's own tool adds it.

What this holds the feature to (docs/transport-plan.md §4):

  * it reproduces CAAT's published figures exactly
  * resupply days and trucks change the truck load and nothing else
  * the footprint and the kg CO2e per kg are untouched by it
  * month by month: a window across two routes uses each for its own months, an open end
    keeps going, a month with no route carries nothing and is named
  * nonsense is refused, overlapping months are refused, nothing is half-saved
  * removing keeps the record and takes the transport out of the figures
  * it reaches the client only on publish, and only their own route
  * the Excel says it, beside the footprint and never in it

    python test_transport.py        (needs the catalogue API running)
"""
import io
import json
import os
import sys
import tempfile

_HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, _HERE)

_TMP = tempfile.mkdtemp(prefix="pp_transport_")
os.environ["MIST_APP_DB"] = os.path.join(_TMP, "transport.db")
os.environ["MIST_TENANT"] = "alpha"
os.environ["MIST_SECRET_KEY"] = "test-only-key-not-a-real-secret"

from fastapi.testclient import TestClient   # noqa: E402
from openpyxl import load_workbook          # noqa: E402

import analysis    # noqa: E402
import auth        # noqa: E402
import catalogue   # noqa: E402
import db          # noqa: E402
import main        # noqa: E402
import publish     # noqa: E402
import store       # noqa: E402
import transport   # noqa: E402

FAILED = []


def P(ok, msg):
    print(("  PASS " if ok else "  FAIL ") + msg)
    if not ok:
        FAILED.append(msg)


TUD = dict(wholesaler="Sligro", dc_location="Berkel & Rodenrijs", km_dc="12", km_campus="5",
           resupply_days="200", trucks_per_day="1", ef_kg_per_tkm="0.363",
           ef_source="CAAT", from_period="2025-01", to_period="", reason="CAAT report p.18")

print("=== it is CAAT's method, to the kilogram ===")
P(abs(transport.per_tonne(TUD) - 10.527) < 1e-9,
  f"TU Delft's settings come to 10.527 kg CO2e per tonne delivered ({transport.per_tonne(TUD)})")
for tonnes, want, what in ((150.777, 1587, "the Sep 2024 - Feb 2025 report"),
                           (640.0515, 6737.8, "the 2025-26 workbook")):
    got = transport.per_tonne(TUD) * tonnes
    P(abs(got - want) <= 1, f"{tonnes} t gives {got:,.1f} kg, as {what} says ({want:,})")
P(transport.describe(TUD)["annual_km"] == 5800, "and 5,800 km a year, as their form computes")

if catalogue.health() is None:
    print("the catalogue API is not running -- start it and try again")
    sys.exit(2)

db.init()
auth.init()
auth.add_tenant("alpha", "Alpha University", "")
auth.add_tenant("beta", "Beta College", "")
con = db.connect()
for tenant in ("alpha", "beta"):
    store.upsert(con, "product", ["tenant", "artikelnr", "description", "category", "ean",
                                  "ean_he", "first_seen", "last_seen"],
                 [(tenant, "194072", "MEYERIJ VOLLE MELK", "ZUIVEL HOUDBAAR", "8710401996797",
                   "", "2025-01", "2025-12")], conflict=["tenant", "artikelnr"])
    for m in range(1, 13):                                     # 1,000 kg a month, all year
        con.execute("INSERT INTO purchase_line (tenant, year, month, klantnr, restaurant, "
                    "city, artikelnr, aantal, omzet, kg, kg_known, quality, source_upload) "
                    "VALUES (?,2025,?,'K1','CANTEEN','X','194072',100,500,1000,1,'complete',?)",
                    (tenant, m, f"upl-{tenant}"))
    store.upsert(con, "upload",
                 ["id", "tenant", "filename", "stored_path", "adapter", "year", "uploaded_at",
                  "periods", "lines", "products", "spend_eur", "verdict", "report_json",
                  "committed_at", "commit_mode", "commit_note", "selected", "archived_at"],
                 [(f"upl-{tenant}", tenant, "Year 2025.xlsx", "", "sligro", 2025,
                   "2026-01-01T00:00:00", json.dumps([f"2025-{m:02d}" for m in range(1, 13)]),
                   12, 1, 6000.0, "go", json.dumps(dict(verdict="go", findings=[])),
                   None, None, None, 1, None)],
                 conflict=["id"])
con.commit()
con.close()
db.invalidate()
auth.create_user("mist", "admin-password-1", "admin", None, "MiSt")
auth.create_user("alphauser", "alpha-password-1", "client", "alpha", "Alpha University")


def signed_in(u, pw):
    c = TestClient(main.app, follow_redirects=False)
    assert c.post("/login", data={"username": u, "password": pw}).status_code == 303
    return c


M = signed_in("mist", "admin-password-1")
A = signed_in("alphauser", "alpha-password-1")
M.post("/admin/viewing", data={"tenant": "alpha", "back": "/"})


def head(client=M, window="all"):
    return client.get(f"/api/analysis?window={window}").json()


print("\n=== no route, no transport ===")
before = head()
P(before.get("transport") is None and "transport_co2_kg" not in before["headline"],
  "a client with no route set carries no transport and no empty line about it")
P("transport-line" not in M.get("/?window=all").text, "and the dashboard says nothing about it")

print("\n=== a route is set ===")
r = M.post("/transport", data=TUD)
P(r.status_code == 303 and "done=" in r.headers["location"], "MiSt saves TU Delft's route")
after = head()
t = after["transport"]
P(after["headline"]["food_kg"] == 12000, f"(12,000 kg of food counted: {after['headline']['food_kg']})")
P(abs(t["kg"] - 12 * 10.527) < 0.2, f"transport is 12 t x 10.527 = {t['kg']} kg")
P(after["headline"]["transport_co2_kg"] == round(t["kg"])
  and after["headline"]["total_co2_kg_incl_transport"]
      == after["headline"]["co2_kg"] + round(t["kg"]) or
  abs(after["headline"]["total_co2_kg_incl_transport"]
      - (after["headline"]["co2_kg"] + t["kg"])) <= 1,
  "the headline carries transport and the total including it")
P(after["headline"]["co2_kg"] == before["headline"]["co2_kg"]
  and after["headline"]["intensity_kg_co2_per_kg"] == before["headline"]["intensity_kg_co2_per_kg"]
  and after["by_month"] == before["by_month"],
  "the footprint, the kg CO2e per kg and the months are exactly what they were")
P(abs(sum(m["kg"] for m in t["per_month"]) - t["kg"]) < 0.5 and len(t["per_month"]) == 12,
  "the months add up to the total")
s = t["settings_used"][0]
P(abs(s["avg_load_t"] - 12 / 200) < 1e-6 and not s["over_truck_class"],
  f"and the average truck load is 12 t / 200 days = {s['avg_load_t']} t")
dash = M.get("/?window=all").text
P("transport-line" in dash and "incl. transport" in dash and "t transport DC" in dash,
  "the dashboard shows it under the food total")

print("\n=== days and trucks only change the truck load ===")
cid = transport.active("alpha")[0]["id"]
transport.remove("alpha", cid, "mist", "test")
transport.add("alpha", dict(TUD, resupply_days="100", trucks_per_day="3"), "mist")
t2 = head()["transport"]
P(abs(t2["kg"] - t["kg"]) < 0.01, f"100 days x 3 trucks: transport unchanged ({t2['kg']})")
P(abs(t2["settings_used"][0]["avg_load_t"] - 12 / 300) < 1e-6, "the truck load is not")

print("\n=== two routes, an open end, and a month with none ===")
transport.remove("alpha", transport.active("alpha")[0]["id"], "mist", "test")
transport.add("alpha", dict(TUD, from_period="2025-01", to_period="2025-06"), "mist")
transport.add("alpha", dict(TUD, km_dc="30", from_period="2025-08", to_period=""), "mist")
t3 = head()["transport"]
by = {m["period"]: m for m in t3["per_month"]}
# each month is kept to 0.1 kg
P(abs(by["2025-03"]["kg"] - 10.527) <= 0.05 and abs(by["2025-09"]["kg"] - 65 * 0.363) <= 0.05,
  "each month uses the route in force for it")
P(t3["months_uncovered"] == ["2025-07"] and by["2025-07"]["kg"] == 0,
  "July has no route: nothing, and named")
cav = [c for c in head()["headline"]["caveats"] if c["code"] == "transport"]
P(len(cav) == 1 and "2025-07" in cav[0]["message"], "and the page says so, once")
P(len(head()["headline"]["caveats"]) == len(head()["headline"]["caveats"]),
  "(serving the figures again does not repeat it)")
P(abs(head(window="AY2024")["transport"]["kg"]
      - sum(by[f"2025-{m:02d}"]["kg"] for m in range(1, 9))) < 0.5,
  "an academic year takes only its own months")

print("\n=== nonsense is refused, and nothing is half-saved ===")
n = len(transport.active("alpha"))
for bad, what in ((dict(TUD, from_period="2025-05"), "overlapping months"),
                  (dict(TUD, from_period="2026-01", reason=""), "no source given"),
                  (dict(TUD, from_period="2026-01", km_dc="0"), "zero distance"),
                  (dict(TUD, from_period="2026-01", km_dc="-3"), "negative distance"),
                  (dict(TUD, from_period="2026-01", ef_kg_per_tkm="0"), "zero factor"),
                  (dict(TUD, from_period="2026-01", to_period="2025-12"), "end before start"),
                  (dict(TUD, from_period="2026-1"), "a month not written as YYYY-MM")):
    r = M.post("/transport", data=bad)
    P(r.status_code == 200 and "not saved" in r.text and len(transport.active("alpha")) == n,
      f"{what} is refused, with the form kept")

print("\n=== it reaches the client only when published, and only theirs ===")
P(A.get("/transport").status_code == 403 and 'href="/transport"' not in A.get("/").text,
  "a client is refused the page and not shown the link")
publish.wait(publish.start("alpha", "mist"), timeout=600)
seen = head(A)
P(seen["transport"] and abs(seen["transport"]["kg"] - t3["kg"]) < 0.01,
  "after publishing the client sees the transport MiSt saw")
transport.remove("alpha", transport.active("alpha")[0]["id"], "mist", "test")
P(abs(head(A)["transport"]["kg"] - t3["kg"]) < 0.01,
  "removing a route changes nothing for the client until the next publish")
P("the transport route changed" in M.get("/").text, "and the admin bar says why the copy is behind")
M.post("/admin/viewing", data={"tenant": "beta", "back": "/"})
P(head().get("transport") is None, "Beta has no route of Alpha's")
M.post("/admin/viewing", data={"tenant": "alpha", "back": "/"})

print("\n=== removing keeps the record ===")
P(any(x["removed_why"] == "test" for x in transport.removed("alpha")),
  "a removed route is kept, with why")
P(transport.digest("alpha") != transport.NONE and len(transport.active("alpha")) == 1,
  "(one route left)")

print("\n=== the Excel says it beside the footprint ===")
wb = load_workbook(io.BytesIO(M.get("/export.xlsx?window=all").content), read_only=True)
summary = {r[0]: r[1] for r in wb["summary"].iter_rows(values_only=True)}
h = head()["headline"]
P(summary.get("transport DC -> client (kg CO2e)") == h["transport_co2_kg"]
  and summary.get("total incl. transport (kg CO2e)") == h["total_co2_kg_incl_transport"]
  and summary["CO2e (kg)"] == h["co2_kg"],
  "the summary has transport and the total, and the CO2e row is still food only")
rows = list(wb["transport"].iter_rows(values_only=True))
total = next(r for r in rows if r and r[0] == "total")
P(abs(total[2] - head()["transport"]["kg"]) < 0.5, "the transport sheet adds up to the headline")
P(any(r and r[0] == "how" for r in rows), "and says how it is worked out")
lines = list(wb["lines"].iter_rows(values_only=True))
P(not any("transport" in str(c).lower() for r in lines[12:] for c in r if c),
  "the lines sheet is untouched: its rows still add up to the food footprint")

import shutil   # noqa: E402
shutil.rmtree(_TMP, ignore_errors=True)
print(f"\n{'ALL PASS' if not FAILED else str(len(FAILED)) + ' FAILED'}")
sys.exit(1 if FAILED else 0)
