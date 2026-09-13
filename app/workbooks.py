"""
The analysis workbook: everything on the dashboard as a spreadsheet, every line on the last tab.

Built in one place because two things build it. /export.xlsx builds it on demand for MiSt,
from the figures as they are right now. Publishing builds it once, from the figures being
published, and keeps the bytes -- so what a client downloads is the same answer as what
their dashboard shows, however far the working figures have moved since.
"""
from __future__ import annotations

import io
import re

import charts
import lines_export


def analysis_workbook(result: dict, scored: dict, client: str,
                      extra_notes: list[str] | None = None) -> bytes:
    """`result` is an analysis; `scored` is catalogue.score_lines() for the same rows."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    head = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor="1A4A2A")

    def sheet(name, cols, rows):
        ws = wb.create_sheet(name[:31])
        ws.append(cols)
        for c in ws[1]:
            c.font, c.fill = head, fill
            c.alignment = Alignment(horizontal="center")
        for r in rows:
            ws.append(r)
        for i, col in enumerate(cols, start=1):
            ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = \
                max(12, min(46, len(str(col)) + 6))
        ws.freeze_panes = "A2"
        return ws

    h = result["headline"]
    wb.remove(wb.active)
    sheet("summary", ["metric", "value"], [
        ["window", h["window"]], ["months", h["months"]],
        ["period from", h["period_from"]], ["period to", h["period_to"]],
        ["food purchased (kg)", h["food_kg"]], ["non-food (kg)", h["nonfood_kg"]],
        ["CO2e (kg)", h["co2_kg"]],
        ["intensity (kg CO2e per kg food)", h["intensity_kg_co2_per_kg"]],
        ["EAT-Lancet score", h["eat_lancet_score"]],
        ["EAT-Lancet profile", h.get("eat_profile")],
        ["spend (EUR)", h["spend_eur"]], ["restaurants", h["restaurants"]],
        ["products", h["products"]], ["lines", h["lines"]],
        ["matched to a specific product (% of weight)",
         h["confidence"]["product_specific_pct_of_weight"]],
        ["per-piece share of spend (%)", h["piece_spend_pct"]],
        # Adjusted and unadjusted side by side, from the same lines as the last tab.
        ["food bought, before adjustments (kg)", round(sum(
            r.get("kg") or 0 for r in scored["rows"] if r.get("is_food")))],
        ["food not counted because of adjustments (kg)", round(sum(
            (r.get("kg") or 0) - (r.get("kg_eff") or 0)
            for r in scored["rows"] if r.get("is_food")))],
    ])
    sheet("caveats", ["severity", "owner", "what you must know"],
          [[c["severity"], c["owner"], c["message"]] for c in h["caveats"]])
    sheet("by month", ["period", "food kg", "kg CO2e", "intensity", "spend EUR", "quality"],
          [[r["period"], r["food_kg"], r["co2_kg"], r["intensity_kg_co2_per_kg"],
            r["spend_eur"], r.get("quality", "")] for r in result["by_month"]])
    sheet("by restaurant", ["restaurant", "food kg", "kg CO2e", "intensity", "spend EUR", "products"],
          [[r["restaurant"], r["food_kg"], r["co2_kg"], r["intensity_kg_co2_per_kg"],
            r["spend_eur"], r["products"]] for r in result["by_restaurant"]])
    sheet("by food group", ["food group", "food kg", "kg CO2e", "% of weight", "% of CO2",
                            "intensity", "products"],
          [[r["food_group"], r["food_kg"], r["co2_kg"], r["pct_of_weight"], r["pct_of_co2"],
            r["intensity_kg_co2_per_kg"], r["products"]] for r in result["by_food_group"]])
    sheet("top contributors", ["artikelnr", "product", "food group", "food kg", "kg CO2e",
                               "kg CO2e per kg", "% of CO2", "precision", "confidence"],
          [[r["artikelnr"], r["description"], r["food_group"], r["food_kg"], r["co2_kg"],
            r["co2_per_kg"], r["pct_of_co2"], charts.grade(r.get("source")),
            r["confidence"]]
           for r in result["top_contributors"]])
    sheet("eat lancet", ["food group", "reference %", "purchased %", "gap"],
          [[r["food_group"], r["reference_pct"], r["purchased_pct"], r["gap_pct"]]
           for r in result["eat_lancet"]["rows"]])
    sheet("data health", ["tier", "what it means", "% of weight", "% of CO2", "products", "specific?"],
          [[t["label"], t["explain"], t["pct_of_weight"], t["pct_of_co2"], t["products"],
            "yes" if t["product_level"] else "no"] for t in result["data_health"]["by_tier"]])
    sheet("zero weight", ["artikelnr", "product", "category", "pack", "unit", "pieces", "spend EUR"],
          [[r["artikelnr"], r["description"], r["category"], r["vp"], r["eenh"],
            r["pieces"], r["spend_eur"]] for r in result["piece_items"]["rows"]])

    # Every line the sheets above are made of. This is the point of the whole download:
    # a reader who doubts a figure can open the rows underneath it here, in the same
    # file, rather than being sent somewhere else to fetch them.
    lines_export.add_sheet(wb, scored["rows"], client, h["window"],
                           version=scored.get("catalogue_version"),
                           extra_notes=extra_notes or [])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def export_name(tenant: str, result: dict) -> str:
    # Only characters every operating system accepts in a file name. A single file's
    # window is labelled "File: ...", and Windows refuses the colon.
    label = re.sub(r"[^A-Za-z0-9()_-]", "", result["headline"]["window"])
    return f"PLANETmeal_{tenant}_{label}.xlsx"
