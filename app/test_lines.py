"""
The per-line export must be every line, and it must add up to the dashboard.

Those are the two things it claims, and both were wrong on the first attempt in ways
that looked fine. The endpoint returned service's `food` frame, which analyse() slices to
food only -- so the export silently dropped 3,123 non-food lines and 183,000 euros of
spend while looking complete. And the kilograms column was the file's weight rather than
the weight the footprint was calculated from, which differ wherever something is sold by
the piece.

Neither was visible without adding the columns up. So this adds them up.

    python test_lines.py        (needs the catalogue API running)
"""
import io
import os
import sys
import tempfile

_HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, _HERE)

_TMP = tempfile.mkdtemp(prefix="pp_lines_")
os.environ["MIST_APP_DB"] = os.path.join(_TMP, "lines.db")
os.environ["MIST_TENANT"] = "acme"
os.environ["MIST_SECRET_KEY"] = "test-only-key"

import openpyxl      # noqa: E402

import auth          # noqa: E402
import catalogue     # noqa: E402
import db            # noqa: E402
import lines_export  # noqa: E402
import selection     # noqa: E402
import store         # noqa: E402

FAILED = []


def P(ok, msg):
    print(("  PASS " if ok else "  FAIL ") + msg)
    if not ok:
        FAILED.append(msg)


def _fresh():
    if not store.IS_POSTGRES:
        return
    con = store.connect()
    for t in ("month_owner", "analysis_run", "upload_line", "upload_product", "upload",
              "purchase_line", "product", "app_user", "tenant"):
        con.execute(f"DROP TABLE IF EXISTS {t} CASCADE")
    con.commit()
    con.close()


_fresh()
db.init()
auth.init()
auth.add_tenant("acme", "Acme University", "")

if catalogue.health() is None:
    print("the catalogue API is not running — start it and try again:")
    print("  python -m uvicorn api:app --port 8077   (in the catalogue repo, src/)")
    sys.exit(2)

# One food line, one non-food line, one sold by the piece. The three cases that made the
# first version wrong: non-food was dropped, and per-piece weight differs from the weight
# the footprint used.
LINES = [
    dict(artikelnr="194072", description="MEYERIJ VOLLE MELK", category="ZUIVEL HOUDBAAR",
         restaurant="Aula", klantnr="K1", year=2025, month=3, aantal=10, omzet=120.0,
         kg=60.0, kg_known=1, ean_ce="8710401996797", ean_he=""),
    dict(artikelnr="119569", description="T.D.SOEPBKR KRTN FSC 250ML 25S",
         category="DISPOSABLES", restaurant="Aula", klantnr="K1", year=2025, month=3,
         aantal=4, omzet=80.0, kg=5.0, kg_known=1, ean_ce="", ean_he=""),
    dict(artikelnr="149475", description="BANAAN BUDGET FAIRTRADE", category="FRUIT VERS",
         restaurant="Aula", klantnr="K1", year=2025, month=4, aantal=100, omzet=180.0,
         kg=0.0, kg_known=0, ean_ce="", ean_he=""),
]

print("=== every line comes back, not only the food ===")
scored = catalogue.score_lines(LINES, label="probe")
rows = scored["rows"]
P(len(rows) == len(LINES), f"{len(rows)} rows out for {len(LINES)} in")
kinds = {r["artikelnr"]: r.get("is_food") for r in rows}
P(set(kinds) == {"194072", "119569", "149475"}, "every article is present")
P(any(v is False for v in kinds.values()),
  "including the non-food one, which the food frame would have dropped")

print("\n=== nutrition does not leak ===")
# named here rather than imported, so the app side asserts the boundary independently
# of the catalogue's own list of what to strip
PY_NUTRITION = ("kcal", "protein_g", "fat_g", "satfat_g", "carb_g", "sugar_g",
                "fibre_g", "sodium_mg", "nutrition_src", "nutrition_conf")
leaked = sorted({k for r in rows for k in r if k in PY_NUTRITION})
P(not leaked, f"no nutrition fields in the export ({leaked or 'none'})")

print("\n=== a line with no weight is present with zeros, not dropped ===")
piece = next(r for r in rows if r["artikelnr"] == "149475")
P(piece is not None, "the per-piece banana line is in the export")
P((piece.get("kg_eff") or 0) == 0 and (piece.get("co2_kg") or 0) == 0,
  "it contributes nothing, and says so rather than vanishing")

print("\n=== the workbook is readable and carries what it promises ===")
data = lines_export.workbook(rows, "Acme University", "probe",
                             version=scored.get("catalogue_version"))
ws = openpyxl.load_workbook(io.BytesIO(data))["lines"]
head_row = next(i for i, r in enumerate(ws.iter_rows(values_only=True), start=1)
                if r[0] == "Period")
heads = [c.value for c in ws[head_row]]
for want in ("Period", "Product", "Kilograms", "Kilograms counted", "kg CO2e",
             "EAT-Lancet group", "Footprint from", "Footprint confidence", "Matched by"):
    P(want in heads, f"column present: {want}")
P(ws.max_row == head_row + len(rows), f"one row per line ({ws.max_row - head_row})")
P(ws.freeze_panes is not None, "the headings stay put when you scroll")

notes = " ".join(str(ws.cell(row=r, column=1).value or "") for r in range(1, head_row))
P("reconcile" in notes.lower(), "the sheet says how to reconcile it with the dashboard")
P("Nutrition is deliberately absent" in notes, "and why nutrition is not in it")

print("\n=== the totals ARE the dashboard's ===")
# The claim the whole export rests on: filter to food, sum two columns, get the headline.
col = {h: i for i, h in enumerate(heads)}
food_kg = co2_food = spend = 0.0
for r in ws.iter_rows(min_row=head_row + 1, values_only=True):
    spend += r[col["Spend EUR"]] or 0
    if r[col["Food?"]] == "food":
        food_kg += r[col["Kilograms counted"]] or 0
        co2_food += r[col["kg CO2e"]] or 0

import analysis  # noqa: E402
con = db.connect()
import json
store.upsert(con, "upload",
             ["id", "tenant", "filename", "stored_path", "adapter", "year", "uploaded_at",
              "periods", "lines", "products", "spend_eur", "verdict", "report_json",
              "committed_at", "commit_mode", "commit_note", "selected", "archived_at"],
             [("f1", "acme", "probe.xlsx", "", "sligro", 2025, "2025-01-01T00:00:00",
               json.dumps(["2025-03", "2025-04"]), 3, 3, 380.0, "go",
               json.dumps(dict(verdict="go", findings=[], summary={}, filename="probe.xlsx",
                               adapter="sligro", lines=3, products=3, spend_eur=380.0,
                               periods=["2025-03", "2025-04"])),
               None, None, None, 1, None)],
             conflict=["id"])
for line in LINES:
    con.execute("INSERT INTO purchase_line (tenant, year, month, klantnr, restaurant, "
                "city, artikelnr, aantal, omzet, kg, kg_known, quality, source_upload) "
                "VALUES ('acme',?,?,?,?,'X',?,?,?,?,?,'complete','f1')",
                (line["year"], line["month"], line["klantnr"], line["restaurant"],
                 line["artikelnr"], line["aantal"], line["omzet"], line["kg"],
                 line["kg_known"]))
    con.execute("INSERT INTO product (tenant, artikelnr, description, category, ean, "
                "ean_he, first_seen, last_seen) VALUES ('acme',?,?,?,?,'','2025-03','2025-04')",
                (line["artikelnr"], line["description"], line["category"], line["ean_ce"]))
con.commit()
con.close()
db.invalidate()

h = analysis.run(window="all", tenant="acme", force=True)["headline"]
P(round(food_kg, 3) == round(h["food_kg"], 3) or abs(food_kg - h["food_kg"]) < 1,
  f"food kilograms match the dashboard ({food_kg:,.1f} vs {h['food_kg']:,})")
P(abs(co2_food - h["co2_kg"]) < 1,
  f"kg CO2e matches the dashboard ({co2_food:,.1f} vs {h['co2_kg']:,})")
P(abs(spend - h["spend_eur"]) < 1,
  f"spend matches, including the non-food ({spend:,.0f} vs {h['spend_eur']:,})")

# --------------------------------------------------------------------------- the two ways in
# The lines are worth nothing to a client who cannot get at them. The first version put
# the per-file download under /upload, which is admin-only, so the one person the sheet
# exists for got a 403 back -- and the window-level download was a separate button that
# produced a file with no context around it. Both are answered by two facts this checks:
# the sheet is the last tab of the dashboard export, and the per-file link is not gated.
from fastapi.testclient import TestClient   # noqa: E402
import main                                 # noqa: E402

auth.create_user("boss", "boss-password-1", "admin", None, "Boss")
auth.create_user("acmeuser", "acme-password-1", "client", "acme", "Acme")


def signed_in(username, password):
    c = TestClient(main.app, follow_redirects=False)
    r = c.post("/login", data={"username": username, "password": password, "next": "/"})
    assert r.status_code == 303, f"could not sign in as {username}"
    return c


client = signed_in("acmeuser", "acme-password-1")

print()
print("=== a client can take the lines away, and cannot change the numbers ===")
r = client.get("/files/f1/lines.xlsx")
P(r.status_code == 200, f"a client may download one file's lines ({r.status_code})")
P(r.headers.get("content-type", "").startswith("application/vnd.openxmlformats"),
  "and gets a workbook, not a login page")

r = client.post("/files/owner", data=dict(year=2025, month=3, upload_id="f1"))
P(r.status_code == 403,
  f"but may NOT name the file that supplies a month ({r.status_code})")

print()
print("=== the dashboard export carries the lines, and they add up to its own summary ===")
r = client.get("/export.xlsx?window=all")
P(r.status_code == 200, f"the export downloads ({r.status_code})")
ex = openpyxl.load_workbook(io.BytesIO(r.content))
P("lines" in ex.sheetnames, f"it has a lines tab ({ex.sheetnames})")
P(ex.sheetnames[-1] == "lines", "last, after the summaries it is the evidence for")

lw = ex["lines"]
lrows = list(lw.iter_rows(values_only=True))
hi = next(i for i, row in enumerate(lrows) if row and row[0] == "Period")
lcol = {h: i for i, h in enumerate(lrows[hi])}
P(len(lrows) - hi - 1 == len(LINES),
  f"one row per purchase line ({len(lrows) - hi - 1})")

x_kg = sum((row[lcol["Kilograms counted"]] or 0)
           for row in lrows[hi + 1:] if row[lcol["Food?"]] == "food")
x_co2 = sum((row[lcol["kg CO2e"]] or 0)
            for row in lrows[hi + 1:] if row[lcol["Food?"]] == "food")
x_spend = sum((row[lcol["Spend EUR"]] or 0) for row in lrows[hi + 1:])

# Checked against the SUMMARY TAB OF THE SAME FILE, not against the dashboard. Two tabs
# of one workbook disagreeing is the failure a reader finds first and forgives least.
said = {row[0]: row[1] for row in ex["summary"].iter_rows(values_only=True)}
P(abs(x_kg - said["food purchased (kg)"]) < 1,
  f"food kilograms match its own summary tab ({x_kg:,.1f} vs {said['food purchased (kg)']:,})")
P(abs(x_co2 - said["CO2e (kg)"]) < 1,
  f"kg CO2e match its own summary tab ({x_co2:,.1f} vs {said['CO2e (kg)']:,})")
P(abs(x_spend - said["spend (EUR)"]) < 1,
  f"spend matches, non-food included ({x_spend:,.0f} vs {said['spend (EUR)']:,})")

print()
print("ALL PASS" if not FAILED else f"{len(FAILED)} FAILED")
sys.exit(1 if FAILED else 0)
