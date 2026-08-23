"""
MiSt template adapter — the shape we hand a client who is not exporting from Sligro.

There are only two honest ways to onboard a new client: write an adapter for whatever
their wholesaler produces, or give them one sheet to fill in. A large client with a fixed
export gets an adapter; everyone else gets this.

Named columns, order irrelevant, extra columns ignored. One row per product per period, so
one month or fifty is the same shape.
"""
from __future__ import annotations

import os
import re

import openpyxl
from openpyxl.utils import get_column_letter

from .sligro import mass_kg, parse_ivp

NAME = "mist_template"
LABEL = "MiSt template"
DESCRIPTION = ("Our standard sheet. Named columns, one row per product per period. "
               "One month or many without changing shape.")

REQUIRED = ["klantnr", "restaurant", "artikelnr", "omschrijving", "artikelgroep",
            "ivp", "vp", "maat", "eenh", "periode", "aantal"]
# ean_ce is not formally required, because a client may not have it — but it is the most
# valuable column in the sheet, so the template says so and the pre-flight reports coverage.
OPTIONAL = ["ean_ce", "ean_he", "brand", "city", "supplier", "omzet"]
COLUMNS = ["klantnr", "restaurant", "city", "artikelnr", "omschrijving", "brand",
           "artikelgroep", "ivp", "vp", "maat", "eenh", "ean_ce", "ean_he", "supplier",
           "periode", "aantal", "omzet"]

# A bare "ean" column is accepted as the consumer-unit barcode, so sheets built against the
# first version of this template keep working.
ALIASES = {
    "ean_ce": ["eance", "ean"],
    "ean_he": ["eanhe"],
    "brand": ["merknaam"],
    "city": ["woonplaats"],
    "supplier": ["leverancier", "leveranciernaam"],
}

HELP = {
    "klantnr": "Ordering account number.",
    "restaurant": "Readable restaurant name.",
    "artikelnr": "Supplier article number. The product's identity.",
    "omschrijving": "Product description as the supplier writes it.",
    "artikelgroep": "Supplier category.",
    "ivp": "Units per pack (Inhoud Verpakking). Must be present — it multiplies the weight.",
    "vp": "Packaging type (Verpakking). 'KG' means the line is sold by weight.",
    "maat": "Pack size. With IVP=1 this is the whole case; with IVP>1 it is per unit.",
    "eenh": "Unit of Maat: KG, GR, LT, CL, ML, DL — or ST for pieces.",
    "periode": "YYYY-MM. One row per product per period.",
    "aantal": "Quantity purchased in that period.",
    "city": "Town, if you have it. Optional.",
    "brand": "Brand name, if you have it. Helps tell similar products apart. Optional.",
    "ean_ce": ("Barcode of the CONSUMER unit — the item itself. The single most useful "
               "column here. An article number belongs to your supplier; a barcode belongs "
               "to the product, so it is what lets us recognise something already worked "
               "out, even from a different wholesaler."),
    "ean_he": ("Barcode of the HANDLING unit — the case or outer. A DIFFERENT number from "
               "ean_ce. Keep them in their own columns; do not merge them."),
    "supplier": "Who supplied it, if more than one. Optional.",
    "omzet": "Spend in EUR. Optional.",
}


class TemplateError(Exception):
    pass


def _norm(v) -> str:
    return re.sub(r"[^a-z0-9]", "", str(v).strip().lower()) if v is not None else ""


def _header(ws):
    """The header row. Must carry 'periode' — that is what distinguishes this template
    from a Sligro export, which also names a Klantnr and an Artikelnr but lays its periods
    out in columns rather than in a column of its own."""
    for row in ws.iter_rows(min_row=1, max_row=8, values_only=True):
        got = {_norm(v) for v in row if v is not None}
        if "periode" not in got or "artikelnr" not in got:
            continue
        if sum(1 for c in REQUIRED if _norm(c) in got) >= 6:
            return {_norm(v): i for i, v in enumerate(row) if v is not None}
    return None


def _period(v):
    """'2025-03', '2025/3', or a real date -> (year, month)."""
    if v is None:
        raise ValueError("empty periode")
    if hasattr(v, "year") and hasattr(v, "month"):
        return int(v.year), int(v.month)
    m = re.match(r"^\s*(\d{4})[-/. ](\d{1,2})\s*$", str(v))
    if not m:
        raise ValueError(f"expected YYYY-MM, got {v!r}")
    y, mo = int(m.group(1)), int(m.group(2))
    if not 1 <= mo <= 12:
        raise ValueError(f"month out of range in {v!r}")
    return y, mo


def detect(path: str, filename: str | None = None) -> bool:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        return _header(wb[wb.sheetnames[0]]) is not None
    finally:
        wb.close()


def read(path: str, filename: str | None = None, year: int | None = None):
    from . import Reading

    filename = filename or os.path.basename(path)
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    idx = _header(ws)
    if idx is None:
        wb.close()
        raise TemplateError("no header row found — this does not match the MiSt template")
    missing = [c for c in REQUIRED if _norm(c) not in idx]
    if missing:
        wb.close()
        raise TemplateError(f"the template is missing required columns: {', '.join(missing)}")

    def cell(row, col):
        i = idx.get(_norm(col))
        if i is None:
            for alt in ALIASES.get(col, []):
                i = idx.get(_norm(alt))
                if i is not None:
                    break
        return row[i] if i is not None and i < len(row) else None

    prod, lines, problems, source_rows = {}, [], [], []
    started = False
    for n, row in enumerate(ws.iter_rows(values_only=True), start=1):
        if not started:                                   # skip down to past the header
            got = {_norm(v) for v in row if v is not None}
            if sum(1 for c in REQUIRED if _norm(c) in got) >= 6:
                started = True
            continue
        if all(v is None for v in row):
            continue
        art = cell(row, "artikelnr")
        if art is None:
            continue
        art = str(art).strip()
        # No quantity means this is not a purchase line — a note, a marker row, a spacer.
        # Checked FIRST so those are never reported as unreadable data.
        a = cell(row, "aantal")
        if not isinstance(a, (int, float)) or a == 0:
            continue
        try:
            y, m = _period(cell(row, "periode"))
        except ValueError as e:
            col = idx.get(_norm("periode"))
            problems.append(dict(
                row=n, artikelnr=art, problem=str(e),
                cell=f"{get_column_letter(col + 1)}{n}" if col is not None else f"row {n}"))
            continue

        maat = cell(row, "maat")
        maat = maat if isinstance(maat, (int, float)) else None
        eenh = str(cell(row, "eenh") or "").strip().upper()
        vp = str(cell(row, "vp") or "").strip().upper()
        ivp = parse_ivp(cell(row, "ivp"))
        ce = str(cell(row, "ean_ce") or "").strip()
        he = str(cell(row, "ean_he") or "").strip()
        o = cell(row, "omzet")

        if art not in prod:
            prod[art] = (art, str(cell(row, "omschrijving") or "").strip(),
                         str(cell(row, "brand") or "").strip(),
                         str(cell(row, "artikelgroep") or "").strip(),
                         str(cell(row, "ivp") or "").strip(), vp, maat, eenh,
                         ce if ce.isdigit() and ce != "0" else "",
                         he if he.isdigit() and he != "0" else "", "")

        kg, known = mass_kg(a, ivp, maat, eenh, vp)
        lines.append((y, m, str(cell(row, "klantnr") or "").strip(),
                      str(cell(row, "restaurant") or "").strip(), "", art, float(a),
                      float(o) if isinstance(o, (int, float)) else 0.0, kg, known, "complete"))
        source_rows.append(n)
    wb.close()

    if not lines:
        raise TemplateError("the template contains no purchase lines with a quantity")

    return Reading(adapter=NAME, filename=filename, products=prod, lines=lines,
                   row_problems=problems, source_rows=source_rows,
                   notes=["Read as the MiSt template; periods come from the periode column."])


# --------------------------------------------------------------------------- the blank sheet
def write_template(path: str) -> str:
    """Write the blank template a client fills in."""
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "purchases"
    head = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor="1A4A2A")          # --ink
    soft = PatternFill("solid", fgColor="E8F0EA")          # --linen

    for i, col in enumerate(COLUMNS, start=1):
        c = ws.cell(row=1, column=i, value=col)
        c.font, c.fill = head, fill
        c.alignment = Alignment(horizontal="center")
        ws.cell(row=2, column=i,
                value="required" if col in REQUIRED else "optional").fill = soft
        ws.column_dimensions[c.column_letter].width = max(12, len(col) + 4)
    ws.freeze_panes = "A3"

    example = {"klantnr": "132201", "restaurant": "APPEL TU DELFT 3ME", "city": "DELFT",
               "artikelnr": "340515", "omschrijving": "KERN KROKET 20% VLEES 28X80G",
               "brand": "KERN", "artikelgroep": "SNACKS",
               "ivp": 1, "vp": "DS", "maat": 2.24, "eenh": "KG",
               "ean_ce": "8712800121619", "ean_he": "18712800121616",
               "supplier": "SLIGRO", "periode": "2025-03", "aantal": 12, "omzet": 214.80}
    for i, col in enumerate(COLUMNS, start=1):
        ws.cell(row=3, column=i, value=example.get(col))

    notes = wb.create_sheet("how to fill this in")
    notes.column_dimensions["A"].width = 18
    notes.column_dimensions["B"].width = 96
    notes.append(["column", "what it means"])
    notes["A1"].font, notes["B1"].font = head, head
    notes["A1"].fill, notes["B1"].fill = fill, fill
    for col in COLUMNS:
        notes.append([col, HELP[col]])
    for extra in [
        ["", ""],
        ["one month or many", "Add a row per product per period. The shape never changes."],
        ["row 3", "An example. Delete it before sending the file back."],
        ["weights", "IVP, Maat and Eenh together give the kilograms. Without them a line "
                    "cannot be weighed and contributes nothing to the footprint."],
        ["pieces", "Eenh = ST with VP not KG means the line is sold per piece and has no "
                   "weight. Those are reported separately, never silently zeroed."],
        ["two barcodes", "ean_ce and ean_he are DIFFERENT numbers - the item and the case. "
                         "Keep them in their own columns. If you only have one, put it in "
                         "ean_ce and leave ean_he empty."],
        ["why ean_ce matters", "An article number is your supplier's. A barcode is the "
                               "product's. With the barcode we can recognise a product we "
                               "have already worked out, even from a different wholesaler, "
                               "instead of guessing from its description."],
    ]:
        notes.append(extra)
    for row in notes.iter_rows(min_row=2):
        row[1].alignment = Alignment(wrap_text=True, vertical="top")

    wb.save(path)
    return path
