"""
Product adjustments: how much of a product counts, for one client, over stated months.

What this replaces was a catalogue rule -- any name containing FRITUUR, 10%, every client,
all time -- which counted TU Delft's "KERN FR.VET VLB NEUTRAAL 10L" in full because the
name did not match. So this checks, with frying oil sold under both kinds of name:

  * nothing is adjusted until MiSt says so, for any client
  * an adjustment applies by article number, to its months only, for its client only
  * the dashboard names it once per label, and only where it changed something
  * the Excel shows what was bought, the share, and what counted, per line and in total
  * figures are recalculated when an adjustment is added or removed, never served stale
  * a client's published figures move only when MiSt publishes, and the bar says why
  * mistakes are refused with a reason: overlaps, bad months, 100%, no reason given
  * removing one keeps the record
  * a client can neither see nor change any of it

    python test_adjustments.py        (needs the catalogue API running)
"""
import io
import json
import os
import sys
import tempfile

_HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, _HERE)

_TMP = tempfile.mkdtemp(prefix="pp_adjust_")
os.environ["MIST_APP_DB"] = os.path.join(_TMP, "adjust_test.db")
os.environ["MIST_TENANT"] = "alpha"
os.environ["MIST_SECRET_KEY"] = "test-only-key-not-a-real-secret"

from fastapi.testclient import TestClient   # noqa: E402
from openpyxl import load_workbook          # noqa: E402

import adjustments  # noqa: E402
import auth         # noqa: E402
import catalogue    # noqa: E402
import db           # noqa: E402
import main         # noqa: E402
import publish      # noqa: E402
import store        # noqa: E402

FAILED = []


def P(ok, msg):
    print(("  PASS " if ok else "  FAIL ") + msg)
    if not ok:
        FAILED.append(msg)


if catalogue.health() is None:
    print("the catalogue API is not running -- start it and try again")
    sys.exit(2)

db.init()
auth.init()
auth.add_tenant("alpha", "Alpha University", "")
auth.add_tenant("beta", "Beta College", "")

OMEGA, FRVET, MILK = "496206", "231708", "194072"
PRODUCTS = [(OMEGA, "KERN FRITUUROLIE OMEGA", "OLIE EN VET", "8711111000001"),
            (FRVET, "KERN FR.VET VLB NEUTRAAL 10L", "OLIE EN VET", "8711111000002"),
            (MILK, "MEYERIJ VOLLE MELK", "ZUIVEL HOUDBAAR", "8710401996797")]
con = db.connect()
for tenant in ("alpha", "beta"):
    for art, desc, cat, ean in PRODUCTS:
        store.upsert(con, "product", ["tenant", "artikelnr", "description", "category", "ean",
                                      "ean_he", "first_seen", "last_seen"],
                     [(tenant, art, desc, cat, ean, "", "2025-01", "2025-12")],
                     conflict=["tenant", "artikelnr"])
        for m in range(1, 13):
            con.execute("INSERT INTO purchase_line (tenant, year, month, klantnr, restaurant, "
                        "city, artikelnr, aantal, omzet, kg, kg_known, quality, source_upload) "
                        "VALUES (?,2025,?,'K1','CANTEEN','X',?,10,50,100,1,'complete',?)",
                        (tenant, m, art, f"upl-{tenant}"))
    store.upsert(con, "upload",
                 ["id", "tenant", "filename", "stored_path", "adapter", "year", "uploaded_at",
                  "periods", "lines", "products", "spend_eur", "verdict", "report_json",
                  "committed_at", "commit_mode", "commit_note", "selected", "archived_at"],
                 [(f"upl-{tenant}", tenant, "Year 2025.xlsx", "", "sligro", 2025,
                   "2026-01-01T00:00:00", json.dumps([f"2025-{m:02d}" for m in range(1, 13)]),
                   36, 3, 1800.0, "go", json.dumps(dict(verdict="go", findings=[])),
                   None, None, None, 1, None)],
                 conflict=["id"])
con.commit()
con.close()
db.invalidate()

auth.create_user("mist", "admin-password-1", "admin", None, "MiSt")
auth.create_user("alphauser", "alpha-password-1", "client", "alpha", "Alpha University")


def signed_in(u, pw):
    c = TestClient(main.app, follow_redirects=False)
    assert c.post("/login", data={"username": u, "password": pw}).status_code == 303
    return c


def kg(client, window="FY2025"):
    r = client.get(f"/api/analysis?window={window}")
    return r.json()["headline"]["food_kg"] if r.status_code == 200 else None


def notes(client):
    r = client.get("/api/analysis?window=FY2025").json()
    return [c["message"] for c in r["headline"]["caveats"] if c["code"] == "adjustment"]


M = signed_in("mist", "admin-password-1")
A = signed_in("alphauser", "alpha-password-1")
M.post("/admin/viewing", data={"tenant": "alpha", "back": "/"})

print("=== nothing is adjusted until MiSt says so ===")
P(kg(M) == 3600, f"three products x 12 months x 100 kg count in full ({kg(M)} kg)")
P(notes(M) == [], "and no adjustment is mentioned")
P(publish.get(publish.start("alpha", "mist", wait=True))["status"] == "live",
  "(Alpha's figures are published as they are)")
# A copy published before adjustments existed carries no record of them. It must not read
# as "the product adjustments changed" for a client that has none.
_live = publish.live("alpha")
_parts = {k: v for k, v in _live["parts"].items() if k != "adjustments"}
_c = db.connect()
_c.execute("UPDATE publication SET parts_json=? WHERE id=?", (json.dumps(_parts), _live["id"]))
_c.commit()
_c.close()
P("They match what you see." in M.get("/").text,
  "a copy published before adjustments existed still matches, when there are none")

print("\n=== finding the products ===")
by_name = {p["artikelnr"] for p in adjustments.products("alpha", "frituur")}
P(by_name == {OMEGA}, "a name search finds the oil whose name says so...")
by_cat = {p["artikelnr"] for p in adjustments.products("alpha", "olie en vet")}
P(by_cat == {OMEGA, FRVET}, "...and a category search finds the one whose name does not")
P([p["artikelnr"] for p in adjustments.products("alpha", FRVET)] == [FRVET], "by article number")
P([p["artikelnr"] for p in adjustments.products("alpha", "8711111000002")] == [FRVET], "by barcode")
hit = adjustments.products("alpha", FRVET)[0]
P(hit["kg"] == 1200 and hit["first"] == "2025-01" and hit["last"] == "2025-12",
  "each with the kilograms bought and the months")
page = M.get("/adjustments?q=olie+en+vet").text
P("KERN FR.VET VLB NEUTRAAL 10L" in page and 'name="artikelnr"' in page,
  "the page lists them with a box to tick")

print("\n=== add one: both oils, 10%, from July ===")
r = M.post("/adjustments", data=dict(artikelnr=[OMEGA, FRVET], share="10", label="frying oil",
                                     from_period="2025-07", to_period="",
                                     reason="used oil collected; agreed with Alpha", q="olie"))
P(r.status_code == 303 and "done=" in r.headers["location"], "saved")
# milk 1200 + oil Jan-Jun in full (2 x 600) + oil Jul-Dec at 10% (2 x 60)
P(kg(M) == 2520, f"the working figure is recalculated with it: 1200 + 1200 + 120 = 2520 ({kg(M)} kg)")
P(notes(M) == ["Only 10% of purchased frying oil is counted."],
  f"the dashboard says so once, under its label, not once per product ({notes(M)})")
dash = M.get("/").text
P("<strong>Adjusted:</strong> Only 10% of purchased frying oil is counted." in dash,
  "as the Adjusted line under the headline figures")
P("agreed with Alpha" not in dash, "and the reason stays MiSt's")
P(kg(M, "files:upl-alpha") == 2520, "a file chosen on its own is counted the same way")

print("\n=== only where it changed something ===")
h1 = db.lines_for(2025, 1, 2025, 6, "alpha")
P(all(l["share"] == 1.0 for l in h1) and adjustments.notes(h1, "alpha") == [],
  "January to June: every line in full, and nothing to say")
h2 = db.lines_for(2025, 7, 2025, 12, "alpha")
oil = [l for l in h2 if l["artikelnr"] in (OMEGA, FRVET)]
P(oil and all(l["share"] == 0.1 and l["adjustment"] == "frying oil" for l in oil)
  and all(l["share"] == 1.0 for l in h2 if l["artikelnr"] == MILK),
  "July to December: the oil at 10%, the milk untouched")

print("\n=== one client's adjustment is not another's ===")
M.post("/admin/viewing", data={"tenant": "beta", "back": "/"})
P(kg(M) == 3600 and notes(M) == [], "Beta, buying the same oils, still counts them in full")
M.post("/admin/viewing", data={"tenant": "alpha", "back": "/"})

print("\n=== the Excel shows bought, share and counted ===")
wb = load_workbook(io.BytesIO(M.get("/export.xlsx?window=FY2025").content), read_only=True)
summary = {r[0]: r[1] for r in wb["summary"].iter_rows(values_only=True)}
P(summary["food purchased (kg)"] == 2520, "the summary total is the adjusted figure")
P(summary["food bought, before adjustments (kg)"] == 3600
  and summary["food not counted because of adjustments (kg)"] == 1080,
  "beside what was bought and what the adjustments took out (3600, 1080)")
P("Only 10% of purchased frying oil is counted." in
  [r[2] for r in wb["caveats"].iter_rows(values_only=True)], "the caveats sheet names it")
rows = list(wb["lines"].iter_rows(values_only=True))
hi = next(i for i, r in enumerate(rows) if r and r[0] == "Period")
col = {h: i for i, h in enumerate(rows[hi])}
aug = next(r for r in rows[hi + 1:] if r[col["Article"]] == FRVET and r[col["Period"]] == "2025-08")
mar = next(r for r in rows[hi + 1:] if r[col["Article"]] == FRVET and r[col["Period"]] == "2025-03")
P(aug[col["Kilograms"]] == 100 and aug[col["Share counted"]] == 0.1
  and abs(aug[col["Kilograms counted"]] - 10) < 1e-9 and aug[col["Adjustment"]] == "frying oil",
  "an August line: 100 kg bought, 10% counted, 10 kg counted, 'frying oil'")
P(mar[col["Share counted"]] == 1 and not mar[col["Adjustment"]], "a March line: in full, no adjustment")

print("\n=== the client's copy moves only on publish ===")
P(kg(A) == 3600 and not notes(A), "Alpha still sees 3600 kg and no adjustment")
P("the product adjustments changed" in M.get("/").text, "the admin bar says why the copy is behind")
publish.start("alpha", "mist", wait=True)
P(kg(A) == 2520 and notes(A) == ["Only 10% of purchased frying oil is counted."],
  "after publishing, Alpha sees 2520 kg and the one line")
P("agreed with Alpha" not in A.get("/").text, "and never the reason")

print("\n=== mistakes are refused, with a reason, and nothing half-saved ===")
before = len(adjustments.active("alpha"))


def refused(**over):
    data = dict(artikelnr=[MILK], share="50", label="milk", from_period="2025-01",
                to_period="2025-03", reason="test", q="melk")
    data.update(over)
    r = M.post("/adjustments", data=data)
    return r.status_code == 200 and "not saved" in r.text, r.text


ok, text = refused(artikelnr=[MILK, OMEGA], from_period="2025-12", to_period="")
P(ok and "already has 10% counted from 2025-07" in text,
  "an overlap with an adjustment in force is refused, naming it")
P(len(adjustments.active("alpha")) == before, "and the milk in the same request was not saved either")
P("value=\"test\"" in text and "checked" in text, "what was typed and ticked is still there")
P(refused(to_period="2024-12")[0], "a last month before the first")
P(refused(from_period="2025-13")[0], "a month that does not exist")
P(refused(share="100")[0], "100% -- that is no adjustment")
P(refused(share="-5")[0] and refused(share="lots")[0], "a negative share, or not a number")
P(refused(reason="  ")[0], "no reason given")
P(refused(artikelnr=["999999"])[0], "an article this client never bought")
P(refused(artikelnr=[])[0], "no product ticked")
P(refused(artikelnr=[OMEGA], from_period="2025-01", to_period="2025-06")[0] is False,
  "(months that do not overlap are fine: the same oil at 50% for January to June)")
P(kg(M) == 2520 - 300, f"and that one is counted too (-300 kg: {kg(M)})")
# Two adjustment lines now apply. On 16 Sep 2026 nine frying oils, each under its own
# name, filled the first page with nine of these; more than one is folded into one line.
dash = M.get("/").text
P("<strong>Adjusted:</strong> 2 purchases count only in part." in dash
  and "Show which" in dash and dash.count("<strong>Adjusted:</strong>") == 1,
  "with more than one, the first page shows one folded line, not one line each")
P("<li>Only 10% of purchased frying oil is counted.</li>" in dash
  and "<li>Only 50% of purchased milk is counted.</li>" in dash,
  "and the list inside it names every one")

print("\n=== removing keeps the record ===")
kg_now = kg(M)
fr = next(a for a in adjustments.active("alpha") if a["artikelnr"] == FRVET)
r = M.post(f"/adjustments/{fr['id']}/remove", data={"why": "fat not recycled after all"})
P(r.status_code == 303 and "done=" in r.headers["location"], "removed")
P(kg(M) == kg_now + 540, f"that fat counts in full again from July (+540 kg: {kg_now} -> {kg(M)})")
gone = [a for a in adjustments.removed("alpha") if a["id"] == fr["id"]]
P(gone and gone[0]["removed_why"] == "fat not recycled after all" and gone[0]["removed_by"] == "mist",
  "it is kept, stamped with who, when and why")
P("fat not recycled after all" in M.get("/adjustments").text, "and listed under Removed adjustments")
P("error=" in M.post(f"/adjustments/{fr['id']}/remove", data={"why": ""}).headers["location"],
  "removing it twice is refused")

print("\n=== a client can neither see nor change any of it ===")
P(A.get("/adjustments").status_code == 403, "the page is refused")
P(A.post("/adjustments", data=dict(artikelnr=[MILK], share="0", label="x", from_period="2025-01",
                                   reason="x")).status_code == 403, "adding is refused")
P(A.post(f"/adjustments/{fr['id']}/remove").status_code == 403, "removing is refused")
P('href="/adjustments"' not in A.get("/").text, "and there is no link to it")
P('href="/adjustments"' in M.get("/").text, "(MiSt has one)")

print(f"\n{'ALL PASS' if not FAILED else str(len(FAILED)) + ' FAILED'}")
sys.exit(1 if FAILED else 0)
